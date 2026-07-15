#!/usr/bin/env python3
"""Force report styling onto a pandoc-generated .docx.

One font everywhere (Latin + CJK), colored headings, and colored-header /
banded / bordered tables. Values are written directly onto the elements
(w:rFonts, w:shd, run color) instead of relying on style inheritance or a
reference document, so Word renders them verbatim across versions.

Usage:
    python style_docx.py FILE.docx [--font "Microsoft JhengHei"]
        [--heading 1A3A5C] [--heading3 2C5378]
        [--header-fill 1A3A5C] [--band F2F5F8] [--border D0D7DE]
"""
import argparse

from docx import Document
from docx.shared import RGBColor
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


def hexrgb(s):
    return RGBColor(int(s[0:2], 16), int(s[2:4], 16), int(s[4:6], 16))


def set_run_font(run, name):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn('w:rFonts'))
    if rf is None:
        rf = OxmlElement('w:rFonts')
        rpr.insert(0, rf)
    for a in ('w:ascii', 'w:hAnsi', 'w:eastAsia', 'w:cs'):
        rf.set(qn(a), name)


def shade(cell, fill):
    tcPr = cell._element.get_or_add_tcPr()
    shd = tcPr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd')
        # w:shd must precede w:noWrap / tcMar / vAlign / ... in CT_TcPr.
        tcPr.insert_element_before(
            shd, 'w:noWrap', 'w:tcMar', 'w:textDirection',
            'w:tcFit', 'w:vAlign', 'w:hideMark')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill)


def set_borders(table, color):
    tblPr = table._element.tblPr
    old = tblPr.find(qn('w:tblBorders'))
    if old is not None:
        tblPr.remove(old)
    b = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        e = OxmlElement('w:' + edge)
        e.set(qn('w:val'), 'single')
        e.set(qn('w:sz'), '4')
        e.set(qn('w:space'), '0')
        e.set(qn('w:color'), color)
        b.append(e)
    # w:tblBorders must precede w:shd / tblLayout / tblCellMar / tblLook in CT_TblPr.
    tblPr.insert_element_before(
        b, 'w:shd', 'w:tblLayout', 'w:tblCellMar',
        'w:tblLook', 'w:tblCaption', 'w:tblDescription')


def set_normal_font(doc, name):
    try:
        n = doc.styles['Normal']
        n.font.name = name
        rpr = n.element.get_or_add_rPr()
        rf = rpr.find(qn('w:rFonts'))
        if rf is None:
            rf = OxmlElement('w:rFonts')
            rpr.insert(0, rf)
        for a in ('w:ascii', 'w:hAnsi', 'w:eastAsia', 'w:cs'):
            rf.set(qn(a), name)
    except Exception as e:
        print("normal style skipped:", e)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("docx")
    ap.add_argument("--font", default="Microsoft JhengHei")
    ap.add_argument("--heading", default="1A3A5C")
    ap.add_argument("--heading3", default="2C5378")
    ap.add_argument("--header-fill", dest="header_fill", default="1A3A5C")
    ap.add_argument("--band", default="F2F5F8")
    ap.add_argument("--border", default="D0D7DE")
    a = ap.parse_args()

    doc = Document(a.docx)
    navy, navy3, white = hexrgb(a.heading), hexrgb(a.heading3), RGBColor(0xFF, 0xFF, 0xFF)

    set_normal_font(doc, a.font)

    for p in doc.paragraphs:
        sn = p.style.name if p.style else ""
        for run in p.runs:
            set_run_font(run, a.font)
        if sn in ("Title", "Heading 1", "Heading 2"):
            for run in p.runs:
                run.font.color.rgb = navy
        elif sn == "Heading 3":
            for run in p.runs:
                run.font.color.rgb = navy3

    for table in doc.tables:
        set_borders(table, a.border)
        for ri, row in enumerate(table.rows):
            for cell in row.cells:
                for p in cell.paragraphs:
                    for run in p.runs:
                        set_run_font(run, a.font)
                if ri == 0:
                    shade(cell, a.header_fill)
                    for p in cell.paragraphs:
                        for run in p.runs:
                            run.font.color.rgb = white
                            run.font.bold = True
                elif ri % 2 == 0:
                    shade(cell, a.band)

    doc.save(a.docx)
    print("styled:", a.docx, "| tables:", len(doc.tables), "| paras:", len(doc.paragraphs))


if __name__ == "__main__":
    main()
